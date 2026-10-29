"""Document parsers that return normalized text records with source metadata."""
from __future__ import annotations

from io import BytesIO, StringIO
from pathlib import Path
import pandas as pd
from pypdf import PdfReader
from docx import Document


def parse_uploaded_file(filename: str, data: bytes) -> list[dict]:
    """Parse one uploaded file into records: {text, source, location}."""
    ext = Path(filename).suffix.lower()
    records: list[dict] = []

    if ext == ".pdf":
        reader = PdfReader(BytesIO(data))
        for page_no, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                records.append({"text": text, "source": filename, "location": f"page {page_no}"})
    elif ext == ".docx":
        doc = Document(BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for i, paragraph in enumerate(paragraphs, start=1):
            records.append({"text": paragraph, "source": filename, "location": f"paragraph {i}"})
        for table_no, table in enumerate(doc.tables, start=1):
            for row_no, row in enumerate(table.rows, start=1):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                line = " | ".join(cells).strip(" |")
                if line:
                    records.append({"text": line, "source": filename,
                                    "location": f"table {table_no}, row {row_no}"})
    elif ext in {".txt", ".md"}:
        text = data.decode("utf-8", errors="replace").strip()
        if text:
            records.append({"text": text, "source": filename, "location": "text"})
    elif ext == ".csv":
        frame = pd.read_csv(BytesIO(data), dtype=str).fillna("")
        for row_no, (_, row) in enumerate(frame.iterrows(), start=2):
            text = " | ".join(f"{col}: {val}" for col, val in row.items() if str(val).strip())
            if text.strip():
                records.append({"text": text, "source": filename, "location": f"row {row_no}"})
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not records:
        raise ValueError(
            f"No extractable text found in {filename}. If this is a scanned PDF, run OCR first."
        )
    return records
